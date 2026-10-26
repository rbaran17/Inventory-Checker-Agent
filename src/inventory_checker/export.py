"""Results export: CSV (one row per sample item) and an Excel workbook for the working papers."""

from __future__ import annotations

import csv
import io
import re
import time
from typing import Any

from .sample import Sample

STATUS_LABEL = {"verified": "Verified", "check": "Check", "exception": "Exception", "not_scanned": "Not scanned"}
RESULT_LABEL = {"match": "Match", "near_match": "Near match - check", "mismatch": "Mismatch", "not_read": "Not read",
                "recorded": "Recorded", "n/a": ""}
DECISION_LABEL = {"accepted": "Accepted", "follow_up": "Follow up", None: ""}


def filename(workspace: str, details: dict[str, Any]) -> str:
    base = details.get("client") or ("demo" if workspace == "demo" else "inventory-count")
    slug = re.sub(r"[^A-Za-z0-9]+", "-", base).strip("-").lower() or "inventory-count"
    return f"{slug}-count-results-{time.strftime('%Y%m%d')}"


def _rows(sample: Sample, assessment: dict[str, Any]) -> tuple[list[str], list[list[Any]]]:
    header = ["Item ID", "Description", "Location", "Expected qty", "Counted qty", "Qty difference", "Status"]
    for name in sample.attribute_names:
        header += [f"{name} (expected)", f"{name} (read)", f"{name} result"]
    header += ["Other identifiers read", "Evidence scan", "Review decision", "Review note"]
    rows = []
    for r in assessment["items"]:
        row = [r["item_id"], r["description"], r["location"], r["expected_qty"], r["counted_qty"],
               r["qty_difference"], STATUS_LABEL[r["status"]]]
        for c in r["checks"]:
            row += [c["expected"], "; ".join(c["observed"]), RESULT_LABEL[c["result"]]]
        row += ["; ".join(f"{o['name']}: {', '.join(o['observed'])}" for o in r["other_attributes"]),
                r["evidence_scan"] or "", DECISION_LABEL[r["review"].get("decision")], r["review"].get("note", "")]
        rows.append(row)
    return header, rows


def results_csv(sample: Sample, assessment: dict[str, Any]) -> str:
    header, rows = _rows(sample, assessment)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(header)
    w.writerows([["" if v is None else v for v in row] for row in rows])
    return buf.getvalue()


def results_xlsx(sample: Sample, assessment: dict[str, Any], details: dict[str, Any], scans: list[dict[str, Any]],
                 ai_label: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_fill, head_font = PatternFill("solid", start_color="1F2A44"), Font(bold=True, color="FFFFFF")
    fills = {"Verified": "DFF3E7", "Check": "FFF1C7", "Exception": "FDE3E3", "Mismatch": "FDE3E3",
             "Near match - check": "FFF1C7", "Match": "DFF3E7", "Not read": "FFF1C7"}

    def table(ws, header: list[str], rows: list[list[Any]], color_cols: tuple[str, ...] = ()) -> None:
        ws.append(header)
        for c in ws[ws.max_row]:
            c.fill, c.font, c.alignment = head_fill, head_font, Alignment(wrap_text=True, vertical="top")
        idx = [i for i, h in enumerate(header) if h in color_cols or h.endswith(" result")]
        for row in rows:
            ws.append(["" if v is None else v for v in row])
            for i in idx:
                cell = ws.cell(ws.max_row, i + 1)
                if (color := fills.get(str(cell.value))):
                    cell.fill = PatternFill("solid", start_color=color)
        for i, h in enumerate(header, 1):
            width = max([len(str(h))] + [len(str(r[i - 1] or "")) for r in rows]) if rows else len(h)
            ws.column_dimensions[get_column_letter(i)].width = min(48, max(10, width + 2))
        ws.freeze_panes = "B2"

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    s = assessment["summary"]
    lines = [("Inventory count results", ""), ("Client", details.get("client", "")),
             ("Location", details.get("location", "")), ("Count date", details.get("count_date", "")),
             ("Counted by", details.get("counter", "")), ("Items to be tested listing", sample.name),
             ("Generated", time.strftime("%Y-%m-%d %H:%M")), ("Detection", ai_label), ("", ""),
             ("Items to be tested", s["total"]), ("Verified", s["verified"]), ("Check", s["check"]),
             ("Exceptions", s["exception"]), ("Not scanned", s["not_scanned"]),
             ("Reviewed (of check + exceptions)", f"{s['reviewed']} of {s['to_review']}"),
             ("Objects seen but not on the listing", len(assessment["not_in_sample"]))]
    for k, v in lines:
        ws.append([k, v])
    ws["A1"].font = Font(bold=True, size=14)
    for row in ws.iter_rows(min_row=2):
        row[0].font = Font(bold=True)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 34, 48

    header, rows = _rows(sample, assessment)
    table(wb.create_sheet("Results"), header, rows, color_cols=("Status",))

    det_rows = []
    for scan in scans:
        for o in scan["objects"]:
            det_rows.append([scan["id"], scan["taken_at"], scan["kind"], o["name"], o["count"],
                             round(o["confidence"], 2), o.get("item_id") or "", o.get("match_basis", ""),
                             "; ".join(f"{a['name']}: {a['value']}" for a in o.get("attributes", [])), o.get("notes", "")])
    table(wb.create_sheet("Detections"), ["Scan", "Time", "Type", "Object", "Count", "Confidence", "Matched item",
                                          "Match basis", "Attributes read", "Notes"], det_rows)
    table(wb.create_sheet("Not on listing"), ["Scan", "Object", "Count", "Attributes read"],
          [[x["scan_id"], x["name"], x["count"], "; ".join(f"{a['name']}: {a['value']}" for a in x["attributes"])]
           for x in assessment["not_in_sample"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
