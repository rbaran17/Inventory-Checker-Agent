"""The items to be tested listing, reconciling what the camera saw with it, and the assessment.

The auditor uploads the items to be tested listing (the audit sample): one row per item to count, with an item ID, a description, the expected
quantity and any attributes to verify (serial number, model number, manufacturer, asset tag, ...).
Every column that is not one of the core columns is an attribute. A blank attribute value means
"record what the camera reads" rather than "verify against this value".
"""

from __future__ import annotations

import csv
import io
import re
from typing import Any

from pydantic import BaseModel, Field

# Core columns and the header names accepted for them (compared case-insensitively, ignoring spaces,
# underscores, dashes, dots, # and slashes).
CORE_COLUMNS: dict[str, tuple[str, ...]] = {
    "item_id": ("itemid", "item", "id", "sku", "itemnumber", "itemno", "partnumber", "partno", "assetid",
                "reference", "ref", "sampleid", "sampleitem"),
    "description": ("description", "itemdescription", "name", "itemname", "desc", "product"),
    "expected_qty": ("expectedqty", "qty", "quantity", "bookqty", "bookquantity", "expectedquantity", "count",
                     "units", "onhand", "systemqty"),
    "location": ("location", "bin", "binlocation", "warehouse", "site", "area", "zone", "shelf"),
}

TEMPLATE_CSV = ("item_id,description,expected_qty,location,serial_number,model_number,manufacturer\n"
                "A-1001,Laptop computer,1,Office 2,C02XK1ZZJG5H,A2337,Apple\n"
                "A-1002,Computer monitor,2,Office 2,,P2422H,Dell\n")


def _key(header: str) -> str:
    return re.sub(r"[\s_\-./#]+", "", header.strip().lower())


def _label(header: str) -> str:
    """Attribute name as shown and sent to the model: 'serial_number' -> 'Serial number'."""
    text = re.sub(r"[_\-]+", " ", header.strip())
    return text[:1].upper() + text[1:]


class SampleItem(BaseModel):
    item_id: str
    description: str
    expected_qty: int = Field(1, ge=0)
    location: str = ""
    attributes: dict[str, str] = Field(default_factory=dict)  # attribute name -> expected value ("" = record only)


class Sample(BaseModel):
    name: str = ""
    attribute_names: list[str] = Field(default_factory=list)
    items: list[SampleItem] = Field(default_factory=list)

    def item(self, item_id: str | None) -> SampleItem | None:
        return next((i for i in self.items if i.item_id == item_id), None)


# ----------------------------------------------------------------------------- parsing


def _rows_from_xlsx(data: bytes) -> list[list[Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - openpyxl is a dependency
        raise ValueError("Excel files need openpyxl installed") from exc
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _rows_from_csv(data: bytes) -> list[list[Any]]:
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return list(csv.reader(io.StringIO(text), dialect))


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def parse_sample(data: bytes, filename: str = "sample.csv") -> Sample:
    """Read a CSV or Excel (.xlsx) sample. Raises ValueError with a message an auditor can act on."""
    is_xlsx = filename.lower().endswith((".xlsx", ".xlsm")) or data[:2] == b"PK"
    rows = _rows_from_xlsx(data) if is_xlsx else _rows_from_csv(data)
    rows = [[_cell(c) for c in r] for r in rows if any(_cell(c) for c in r)]
    if not rows:
        raise ValueError("The file is empty")
    header = rows[0]
    mapping: dict[int, str] = {}
    attributes: dict[int, str] = {}
    for idx, col in enumerate(header):
        if not col:
            continue
        core = next((name for name, aliases in CORE_COLUMNS.items() if _key(col) in aliases), None)
        if core and core not in mapping.values():
            mapping[idx] = core
        else:
            attributes[idx] = _label(col)
    if "description" not in mapping.values():
        raise ValueError(f"No description column found. Columns seen: {', '.join(h for h in header if h)}. "
                         "Name one 'description' (or 'item name').")
    items: list[SampleItem] = []
    seen: set[str] = set()
    for line, row in enumerate(rows[1:], start=2):
        row = row + [""] * (len(header) - len(row))
        core = {name: row[i] for i, name in mapping.items()}
        if not core.get("description") and not core.get("item_id"):
            continue
        item_id = core.get("item_id") or f"ROW-{line}"
        if item_id in seen:
            raise ValueError(f"Row {line}: item ID {item_id} appears twice")
        seen.add(item_id)
        qty_text = core.get("expected_qty", "")
        try:
            qty = int(float(qty_text.replace(",", ""))) if qty_text else 1
        except ValueError as exc:
            raise ValueError(f"Row {line}: expected quantity '{qty_text}' is not a number") from exc
        items.append(SampleItem(item_id=item_id, description=core.get("description") or item_id,
                                expected_qty=max(qty, 0), location=core.get("location", ""),
                                attributes={name: row[i] for i, name in attributes.items()}))
    if not items:
        raise ValueError("The listing has a header row but no items")
    return Sample(name=filename, attribute_names=list(attributes.values()), items=items)


# ----------------------------------------------------------------------------- demo sample

DEMO_ATTRIBUTES = ["Brand", "Model number", "Serial number"]


def _demo(item_id: str, desc: str, qty: int, location: str, brand: str = "", model: str = "",
          serial: str = "") -> SampleItem:
    return SampleItem(item_id=item_id, description=desc, expected_qty=qty, location=location,
                      attributes={"Brand": brand, "Model number": model, "Serial number": serial})


def demo_sample() -> Sample:
    """Everyday home/office items. Attributes are blank (record what is read): edit them in the Demo tab
    to the values on your own items to see matches, or to something else to see exceptions."""
    return Sample(name="Demo: home & office items", attribute_names=list(DEMO_ATTRIBUTES), items=[
        _demo("DEMO-01", "Laptop computer", 1, "Desk"),
        _demo("DEMO-02", "Computer mouse", 1, "Desk"),
        _demo("DEMO-03", "Keyboard", 1, "Desk"),
        _demo("DEMO-04", "Mobile phone", 1, "Desk"),
        _demo("DEMO-05", "Headphones or earbuds", 1, "Desk"),
        _demo("DEMO-06", "Charger or power adapter", 1, "Desk"),
        _demo("DEMO-07", "Coffee mug", 2, "Kitchen"),
        _demo("DEMO-08", "Water bottle", 1, "Kitchen"),
        _demo("DEMO-09", "Book or notebook", 3, "Shelf"),
        _demo("DEMO-10", "Pen or pencil", 4, "Desk"),
        _demo("DEMO-11", "TV or media remote control", 1, "Living room"),
        _demo("DEMO-12", "Television or computer monitor", 1, "Living room"),
    ])


# ----------------------------------------------------------------------------- comparing values


def normalise(value: str | None) -> str:
    """For comparing codes read from labels: uppercase, letters and digits only, common prefixes dropped."""
    v = (value or "").upper()
    v = re.sub(r"^\s*(S/?N|SERIAL(\s*(NO|NUMBER))?|MODEL(\s*(NO|NUMBER))?|P/?N|REF)\s*[:#.]?\s*", "", v)
    return re.sub(r"[^A-Z0-9]", "", v)


def _distance(a: str, b: str) -> int:
    if abs(len(a) - len(b)) > 2:
        return 99
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb)))
        prev = cur
    return prev[-1]


def compare(expected: str, observed: list[str]) -> str:
    """match | near_match | mismatch | not_read | recorded | n/a"""
    exp = normalise(expected)
    obs = [normalise(o) for o in observed if normalise(o)]
    if not exp:
        return "recorded" if obs else "n/a"
    if not obs:
        return "not_read"
    if exp in obs:
        return "match"
    # Short values (brands, short model codes) must match exactly; longer codes tolerate one misread
    # character (O/0, I/1, B/8...) but are flagged for the auditor to check.
    if len(exp) >= 6 and any(_distance(exp, o) <= 1 for o in obs):
        return "near_match"
    return "mismatch"


# ----------------------------------------------------------------------------- reconciliation


def reconcile(sample: Sample, obj: dict[str, Any]) -> tuple[str | None, str]:
    """Which sample item a detected object is. Returns (item_id, basis).

    A code read from the object's label (serial, model, asset tag...) that equals an item's expected value is
    stronger evidence than the model's description-based guess, so it wins.
    """
    observed = {(a.get("name") or "").strip().lower(): normalise(a.get("value")) for a in obj.get("attributes", [])}
    best: tuple[int, str, str] | None = None
    for item in sample.items:
        hits = [name for name, exp in item.attributes.items()
                if len(normalise(exp)) >= 4 and observed.get(name.lower()) == normalise(exp)]
        if hits:
            score = len(hits) + (2 if any("serial" in h.lower() or "asset" in h.lower() for h in hits) else 0)
            if best is None or score > best[0]:
                best = (score, item.item_id, hits[0])
    if best:
        return best[1], best[2]
    guess = obj.get("matched_item_id")
    if sample.item(guess):
        return guess, "description"
    return None, "none"


def assess(sample: Sample, scans: list[dict[str, Any]], reviews: dict[str, dict[str, Any]] | None = None
           ) -> dict[str, Any]:
    """Per sample item: counted quantity, each attribute's result, and an overall status.

    The counted quantity is the highest count in any single scan, so filming a shelf twice does not
    double count. Attribute values are collected across all scans.
    """
    reviews = reviews or {}
    counts: dict[str, dict[str, Any]] = {}
    observed: dict[str, dict[str, list[str]]] = {}
    not_in_sample = []
    for scan in scans:
        per_scan: dict[str, int] = {}
        for obj in scan["objects"]:
            item_id = obj.get("item_id")
            if not sample.item(item_id):
                not_in_sample.append({"scan_id": scan["id"], "name": obj["name"], "count": obj["count"],
                                      "attributes": obj.get("attributes", [])})
                continue
            per_scan[item_id] = per_scan.get(item_id, 0) + int(obj["count"])
            for a in obj.get("attributes", []):
                if (a.get("value") or "").strip():
                    vals = observed.setdefault(item_id, {}).setdefault((a.get("name") or "").strip().lower(), [])
                    if a["value"].strip() not in vals:
                        vals.append(a["value"].strip())
        for item_id, qty in per_scan.items():
            if item_id not in counts or qty > counts[item_id]["qty"]:
                counts[item_id] = {"qty": qty, "scan_id": scan["id"]}
    rows = []
    for item in sample.items:
        seen = counts.get(item.item_id)
        obs = observed.get(item.item_id, {})
        checks = []
        for name in sample.attribute_names:
            values = obs.get(name.lower(), [])
            checks.append({"name": name, "expected": item.attributes.get(name, ""), "observed": values,
                           "result": compare(item.attributes.get(name, ""), values)})
        extra = [{"name": n, "observed": v} for n, v in obs.items()
                 if n not in {a.lower() for a in sample.attribute_names}]
        qty_diff = None if seen is None else seen["qty"] - item.expected_qty
        results = {c["result"] for c in checks}
        if seen is None:
            status = "not_scanned"
        elif qty_diff != 0 or "mismatch" in results:
            status = "exception"
        elif "near_match" in results or "not_read" in results:
            status = "check"
        else:
            status = "verified"
        review = reviews.get(item.item_id, {})
        rows.append({**item.model_dump(), "counted_qty": seen["qty"] if seen else None, "qty_difference": qty_diff,
                     "evidence_scan": seen["scan_id"] if seen else None, "checks": checks, "other_attributes": extra,
                     "status": status, "review": review})
    summary = {s: sum(1 for r in rows if r["status"] == s) for s in ("verified", "check", "exception", "not_scanned")}
    summary["total"] = len(rows)
    summary["reviewed"] = sum(1 for r in rows if r["status"] in ("check", "exception") and r["review"].get("decision"))
    summary["to_review"] = summary["check"] + summary["exception"]
    return {"items": rows, "summary": summary, "not_in_sample": not_in_sample}
