from __future__ import annotations

import asyncio
import io
import json
import zipfile

from agent_framework import Message
from openpyxl import load_workbook

from inventory_checker.agents import build_observer_copilot, draft_memo, read_evidence
from inventory_checker.agents.closeout import CloseoutRequest, run_closeout
from inventory_checker.models import FieldSource, SessionStatus

from .conftest import OBS, ScriptedChatClient, text_turn, tool_turn


def run(coro):
    return asyncio.run(coro)


def test_evidence_reader_sends_image_and_parses_structured_draft():
    extract = {"tag_numbers": ["1043"], "skus": ["FG-2002"], "suggested_count": 48, "count_basis": "6 x 8 face",
               "confidence": 0.82}
    client = ScriptedChatClient([text_turn(json.dumps(extract))])
    out = run(read_evidence(client, b"\xff\xd8jpeg", "image/jpeg", context="zone B"))
    assert out.skus == ["FG-2002"] and out.suggested_count == 48
    sent = client.calls[0]
    assert sent["options"]["response_format"].__name__ == "EvidenceExtract"
    types = [c.type for m in sent["messages"] for c in m.contents]
    assert "data" in types or "uri" in types  # the image went to the model


def test_copilot_write_needs_approval_then_saves_as_accepted(service, session):
    client = ScriptedChatClient([
        tool_turn("record_test_count", {"sku": "FG-2002", "auditor_count": 255, "direction": "floor_to_sheet",
                                        "zone": "B", "tag_no": "1043"}),
        text_turn("Saved the test count."),
    ])
    agent = build_observer_copilot(client, service, session.id, OBS)
    s = agent.create_session()
    resp = run(agent.run("Counted 255 of FG-2002 at tag 1043 in zone B", session=s))
    assert service.records(session.id) == []  # nothing saved before approval
    [req] = resp.user_input_requests
    assert req.function_call.name == "record_test_count"
    resp = run(agent.run(Message(role="user", contents=[req.to_function_approval_response(True)]), session=s))
    [rec] = service.records(session.id)
    assert rec.data["auditor_count"] == 255 and rec.data["difference"] == -5
    assert rec.field_sources["auditor_count"] == FieldSource.ACCEPTED
    assert rec.field_sources["client_count"] == FieldSource.IMPORTED


def test_copilot_rejected_draft_is_not_saved(service, session):
    client = ScriptedChatClient([
        tool_turn("record_cutoff_document", {"type": "receiving", "document_no": "GRN-9", "date": "2026-12-31"}),
        text_turn("Okay, not saved."),
    ])
    agent = build_observer_copilot(client, service, session.id, OBS)
    s = agent.create_session()
    [req] = run(agent.run("last GRN was 9", session=s)).user_input_requests
    run(agent.run(Message(role="user", contents=[req.to_function_approval_response(False)]), session=s))
    assert service.records(session.id) == []


def test_copilot_read_tools_run_without_approval(service, session):
    client = ScriptedChatClient([tool_turn("lookup_listing_item", {"sku": "FG-2003"}), text_turn("18 EA in B-02-01")])
    agent = build_observer_copilot(client, service, session.id, OBS)
    resp = run(agent.run("where is FG-2003?"))
    assert resp.user_input_requests == [] and "18" in resp.text
    tool_results = [c for m in client.calls[1]["messages"] for c in m.contents if c.type == "function_result"]
    assert "B-02-01" in str(tool_results[0].result)


def test_closeout_workflow_blocks_then_exports(service, session, device):
    blocked = run(run_closeout(service, CloseoutRequest(session.id, OBS)))
    assert blocked.status == "blocked" and blocked.open_items
    assert service.get_session(session.id).status == SessionStatus.OPEN

    rec = service.create_record(OBS, session.id, "test_count",
                                {"sku": "FG-2002", "auditor_count": 250, "direction": "floor_to_sheet",
                                 "explanation": "10 units on open sales order staged in dock"})
    from inventory_checker.models import utcnow
    from inventory_checker.store import sha256_hex
    img = b"\xff\xd8photo"
    service.ingest_capture(device, img, claimed_sha256=sha256_hex(img), kind="photo", media_type="image/jpeg",
                           capture_time=utcnow(), record_id=rec.id)
    service.create_record(OBS, session.id, "tag_control", {"range_start": 1000, "range_end": 1200,
                                                           "used": ["1000-1150"], "unused": ["1151-1200"]})
    client = ScriptedChatClient([text_turn("One test count was performed [rec]. Difference explained.")])
    ready = run(run_closeout(service, CloseoutRequest(session.id, OBS), client))
    assert ready.status == "ready"
    assert service.get_session(session.id).status == SessionStatus.CLOSED
    assert "AI draft" in ready.memo_markdown and rec.id in ready.memo_markdown

    wb = load_workbook(io.BytesIO(ready.schedule_xlsx))
    assert {"Summary", "Test Counts", "Evidence Manifest", "Tag Control"} <= set(wb.sheetnames)
    ws = wb["Test Counts"]
    header = [c.value for c in ws[1]]
    row = dict(zip(header, [c.value for c in ws[2]]))
    assert row["sku"] == "FG-2002" and row["difference"] == -10 and row["over_tolerance"] is True
    assert ws.cell(2, 1).fill.start_color.rgb.endswith("FFE699")  # amber

    z = zipfile.ZipFile(io.BytesIO(ready.evidence_package))
    manifest = json.loads(z.read("manifest.json"))
    [item] = manifest["items"]
    assert sha256_hex(z.read(item["file"])) == item["sha256"]
    assert {"audit_trail.json", "records.json", "SHA256SUMS.csv", "observation_memo.md"} <= set(z.namelist())


def test_memo_without_model_is_deterministic(service, session):
    service.create_record(OBS, session.id, "procedure_observation",
                          {"area": "movement_freeze", "compliant": "no", "note": "Forklift moved stock in zone B"})
    memo = run(draft_memo(service, session.id))
    assert "AI draft" not in memo and "Movement freeze: **no**" in memo
