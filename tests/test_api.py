from __future__ import annotations

import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from inventory_checker.api import create_app
from inventory_checker.config import Settings
from inventory_checker.detection import Attribute, DetectedObject, SceneResult

KEY = "test-key-123"
LISTING = (b"item_id,description,expected_qty,location,serial_number,model_number\n"
           b"A-1,Laptop computer,1,Office,C02XK1ZZJG5H,A2337\n"
           b"A-2,Laptop computer,1,Office,C02XK1ZZAAAA,A2337\n"
           b"B-1,Coffee mug,2,Kitchen,,\n")


def make(detector=None):
    app = create_app(Settings(public_base_url="https://tunnel.example", access_key=KEY), detector=detector,
                     use_env_client=False)
    return TestClient(app)


def authed(detector=None):
    c = make(detector)
    r = c.get(f"/?key={KEY}", follow_redirects=False)
    assert r.status_code == 307 and r.headers["location"] == "/"
    return c


def frames(n=1):
    return [("frames", (f"f{i}.jpg", b"\xff\xd8frame", "image/jpeg")) for i in range(n)]


def test_access_key_required():
    c = make()
    assert c.get("/").status_code == 401 and c.get("/api/state").status_code == 401
    assert c.get("/?key=wrong").status_code == 401
    c = authed()
    assert "Inventory Count" in c.get("/").text and "Count Camera" in c.get("/scan").text
    qr = c.get("/api/qr").json()
    assert qr["url"] == f"https://tunnel.example/scan?key={KEY}" and "<svg" in qr["svg"]
    assert c.get("/healthz").json()["status"] == "ok"


def test_engagement_workflow_end_to_end():
    calls = []

    async def detector(imgs, sample):
        calls.append((len(imgs), [i.item_id for i in sample.items]))
        # The AI guesses A-1 for the laptop, but the serial it read belongs to A-2.
        return SceneResult(objects=[
            DetectedObject(name="laptop", count=1, confidence=0.95, matched_item_id="A-1",
                           attributes=[Attribute(name="Serial number", value="C02X-K1ZZ-AAAA"),
                                       Attribute(name="Model number", value="A2338")]),
            DetectedObject(name="coffee mug", count=2, confidence=0.9, matched_item_id="B-1"),
            DetectedObject(name="rubber duck", count=1, confidence=0.6),
        ], image_quality="good")

    c = authed(detector)
    s = c.get("/api/state").json()
    assert s["active"] == "engagement" and s["sample"] is None
    r = c.post("/api/scans", files=frames())
    assert r.status_code == 409 and "items to be tested listing" in r.json()["detail"]

    c.post("/api/engagement/details", json={"client": "Acme Ltd", "location": "Main WH", "count_date": "2026-12-31",
                                            "counter": "R. Auditor"})
    s = c.post("/api/engagement/sample", files={"file": ("listing.csv", LISTING, "text/csv")}).json()
    assert [i["item_id"] for i in s["sample"]["items"]] == ["A-1", "A-2", "B-1"]
    assert s["sample"]["attribute_names"] == ["Serial number", "Model number"]
    assert c.post("/api/engagement/sample", files={"file": ("x.csv", b"foo\n1\n", "text/csv")}).status_code == 400

    r = c.post("/api/scans", files=frames(4)).json()
    assert calls[0] == (4, ["A-1", "A-2", "B-1"]) and r["workspace"] == "engagement"
    laptop = r["scan"]["objects"][0]
    assert laptop["item_id"] == "A-2" and laptop["match_basis"] == "Serial number" and laptop["suggested_item_id"] == "A-1"

    rows = {x["item_id"]: x for x in c.get("/api/state").json()["assessment"]["items"]}
    assert rows["A-2"]["status"] == "exception"  # model number mismatch
    assert {ch["name"]: ch["result"] for ch in rows["A-2"]["checks"]} == {"Serial number": "match",
                                                                          "Model number": "mismatch"}
    assert rows["B-1"]["status"] == "verified" and rows["A-1"]["status"] == "not_scanned"

    s = c.post("/api/items/A-2/review", json={"decision": "follow_up", "note": "Model label differs"}).json()
    assert s["assessment"]["summary"]["reviewed"] == 1
    s = c.post("/api/scans/S001/objects/2", json={"item_id": "B-1"}).json()  # auditor fixes a match
    assert s["scans"][0]["objects"][2]["mapped_by"] == "you"
    assert c.post("/api/scans/S001/objects/2", json={"item_id": "NOPE"}).status_code == 400

    csv_text = c.get("/api/export.csv").text
    assert csv_text.startswith("Item ID,Description") and "Model label differs" in csv_text
    x = c.get("/api/export.xlsx")
    assert "acme-ltd-count-results" in x.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Summary", "Results", "Detections", "Not on listing"]
    assert wb["Summary"]["B2"].value == "Acme Ltd"
    header = [cell.value for cell in wb["Results"][1]]
    assert "Serial number (read)" in header and "Model number result" in header


def test_demo_is_separate_and_editable():
    async def detector(imgs, sample):
        return SceneResult(objects=[DetectedObject(name="mug", count=2, confidence=0.9, matched_item_id="DEMO-07",
                                                   attributes=[Attribute(name="Brand", value="IKEA")])])

    c = authed(detector)
    s = c.post("/api/workspace", json={"name": "demo"}).json()
    assert s["active"] == "demo" and len(s["sample"]["items"]) == 12
    c.post("/api/scans", files=frames())
    row = next(r for r in c.get("/api/state").json()["assessment"]["items"] if r["item_id"] == "DEMO-07")
    assert row["status"] == "verified" and row["checks"][0]["result"] == "recorded"
    s = c.post("/api/demo/items/DEMO-07", json={"attributes": {"Brand": "Muji"}}).json()
    row = next(r for r in s["assessment"]["items"] if r["item_id"] == "DEMO-07")
    assert row["status"] == "exception" and row["checks"][0]["result"] == "mismatch"
    assert c.post("/api/demo/items/NOPE", json={"expected_qty": 1}).status_code == 404

    s = c.post("/api/workspace", json={"name": "engagement"}).json()
    assert s["scans"] == [] and s["workspaces"]["demo"]["scans"] == 1  # demo scans stay in demo
    s = c.post("/api/demo/reset").json()
    assert s["workspaces"]["demo"]["scans"] == 0


def test_without_ai_scans_are_refused_clearly():
    c = authed(None)
    c.post("/api/workspace", json={"name": "demo"})
    r = c.post("/api/scans", files=frames())
    assert r.status_code == 503 and "ANTHROPIC_API_KEY" in r.json()["detail"]


def test_template_download():
    c = authed()
    r = c.get("/api/sample-template.csv")
    assert r.text.startswith("item_id,description,expected_qty") and "attachment" in r.headers["content-disposition"]


def test_access_key_survives_restarts(monkeypatch, tmp_path):
    from inventory_checker.config import load_or_create_access_key

    monkeypatch.delenv("APP_ACCESS_KEY", raising=False)
    monkeypatch.setenv("APP_ACCESS_KEY_FILE", str(tmp_path / "key"))
    first = load_or_create_access_key()
    assert load_or_create_access_key() == first and (tmp_path / "key").stat().st_mode & 0o777 == 0o600
    monkeypatch.setenv("APP_ACCESS_KEY", "fixed")
    assert load_or_create_access_key() == "fixed"


def test_phone_results_view_and_removal():
    async def detector(imgs, sample):
        return SceneResult(objects=[DetectedObject(name="mug", count=1, confidence=0.9, matched_item_id="DEMO-07"),
                                    DetectedObject(name="duck", count=1, confidence=0.5)])

    c = authed(detector)
    c.post("/api/workspace", json={"name": "demo"})
    c.post("/api/scans", files=frames())
    c.post("/api/scans", files=frames(4))
    v = c.get("/api/phone/scans").json()
    assert v["workspace"] == "demo" and [s["id"] for s in v["scans"]] == ["S001", "S002"] and v["summary"]["total"] == 12
    v = c.delete("/api/scans/S002/objects/1").json()  # drop the background object only
    assert [o["name"] for o in v["scans"][1]["objects"]] == ["mug"]
    assert c.delete("/api/scans/S002/objects/5").status_code == 404
    c.delete("/api/scans/S001")
    assert [s["id"] for s in c.get("/api/phone/scans").json()["scans"]] == ["S002"]
